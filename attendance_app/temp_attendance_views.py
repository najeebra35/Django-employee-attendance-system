#temp_attendance_views.py
"""
Temporary / outsource hourly worker management — a lightweight, separate
system from the main Employee/Attendance models. Workers are added quickly
(name + optional rate), attendance is marked only for days actually worked
(present-only), and hours are entered manually per day since breaks etc.
mean simple in/out subtraction isn't accurate. Totals are summed
automatically on export.
"""
import datetime
import io
from collections import defaultdict

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone

from .models import TemporaryEmployee, TemporaryAttendance
from .middleware import log_activity
from .views import has_perm


def _f(request, key, default=None):
    v = request.POST.get(key, '').strip()
    if v == '':
        return default
    try:
        return float(v)
    except ValueError:
        return default


def _d(request, key):
    v = request.POST.get(key, '').strip()
    try:
        return datetime.date.fromisoformat(v)
    except ValueError:
        return None


# ─────────────────────────────────────────────────────────────────────────────
# Temporary Employees (simple CRUD)
# ─────────────────────────────────────────────────────────────────────────────

@login_required
def temp_employee_list(request):
    if not has_perm(request.user, 'temp_attendance_view'):
        messages.error(request, 'Access denied.')
        return redirect('dashboard')

    workers = TemporaryEmployee.objects.all()
    search = request.GET.get('search', '')
    status_filter = request.GET.get('status', '')
    if search:
        workers = workers.filter(name__icontains=search)
    if status_filter:
        workers = workers.filter(status=status_filter)

    return render(request, 'attendance_app/temp_employee_list.html', {
        'workers': workers,
        'search': search,
        'status_filter': status_filter,
        'can_manage': has_perm(request.user, 'temp_attendance_manage'),
    })


@login_required
def temp_employee_add(request):
    if not has_perm(request.user, 'temp_attendance_manage'):
        messages.error(request, 'Access denied.')
        return redirect('temp_employee_list')

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        if not name:
            messages.error(request, 'Name is required.')
            return render(request, 'attendance_app/temp_employee_form.html', {
                'form_data': request.POST, 'action': 'Add',
            })

        w = TemporaryEmployee.objects.create(
            name=name,
            designation=request.POST.get('designation', '').strip() or None,
            mobile=request.POST.get('mobile', '').strip() or None,
            hourly_rate=_f(request, 'hourly_rate'),
            status=request.POST.get('status', 'active'),
            notes=request.POST.get('notes', '').strip() or None,
            created_by=request.user,
        )
        log_activity(request.user, 'create', 'TemporaryEmployee', w, f'Added temp worker {w.name}', request)
        messages.success(request, f'{w.name} added.')
        return redirect('temp_employee_list')

    return render(request, 'attendance_app/temp_employee_form.html', {'action': 'Add', 'form_data': defaultdict(str)})


@login_required
def temp_employee_edit(request, pk):
    if not has_perm(request.user, 'temp_attendance_manage'):
        messages.error(request, 'Access denied.')
        return redirect('temp_employee_list')

    w = get_object_or_404(TemporaryEmployee, pk=pk)
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        if not name:
            messages.error(request, 'Name is required.')
            return render(request, 'attendance_app/temp_employee_form.html', {
                'worker': w, 'action': 'Edit',
            })

        w.name = name
        w.designation = request.POST.get('designation', '').strip() or None
        w.mobile = request.POST.get('mobile', '').strip() or None
        w.hourly_rate = _f(request, 'hourly_rate')
        w.status = request.POST.get('status', 'active')
        w.notes = request.POST.get('notes', '').strip() or None
        w.save()
        log_activity(request.user, 'update', 'TemporaryEmployee', w, f'Updated temp worker {w.name}', request)
        messages.success(request, f'{w.name} updated.')
        return redirect('temp_employee_list')

    return render(request, 'attendance_app/temp_employee_form.html', {'worker': w, 'action': 'Edit', 'form_data': defaultdict(str)})


@login_required
def temp_employee_delete(request, pk):
    if not has_perm(request.user, 'temp_attendance_manage'):
        messages.error(request, 'Access denied.')
        return redirect('temp_employee_list')

    w = get_object_or_404(TemporaryEmployee, pk=pk)
    if request.method == 'POST':
        name = w.name
        w.delete()
        log_activity(request.user, 'delete', 'TemporaryEmployee', description=f'Deleted temp worker {name}', request=request)
        messages.success(request, f'{name} deleted.')
    return redirect('temp_employee_list')


# ─────────────────────────────────────────────────────────────────────────────
# Temporary Attendance — bulk mark by date (only present days get a record)
# ─────────────────────────────────────────────────────────────────────────────

@login_required
def temp_attendance_mark(request):
    if not has_perm(request.user, 'temp_attendance_manage'):
        messages.error(request, 'Access denied.')
        return redirect('temp_attendance_list')

    today = timezone.localdate()
    date_str = request.GET.get('date', today.isoformat())
    try:
        selected_date = datetime.date.fromisoformat(date_str)
    except ValueError:
        selected_date = today

    if request.method == 'POST':
        selected_date = _d(request, 'date') or today
        saved = 0
        cleared = 0
        for w in TemporaryEmployee.objects.filter(status='active'):
            hours_raw = request.POST.get(f'hours_{w.pk}', '').strip()
            notes = request.POST.get(f'notes_{w.pk}', '').strip()
            existing = TemporaryAttendance.objects.filter(temp_employee=w, date=selected_date).first()

            if hours_raw == '':
                # No hours entered — remove any existing record (not worked that day)
                if existing:
                    existing.delete()
                    cleared += 1
                continue

            try:
                hours = float(hours_raw)
            except ValueError:
                continue
            if hours <= 0:
                if existing:
                    existing.delete()
                    cleared += 1
                continue

            if existing:
                existing.hours_worked = hours
                existing.notes = notes
                existing.updated_by = request.user
                existing.save()
            else:
                TemporaryAttendance.objects.create(
                    temp_employee=w, date=selected_date, hours_worked=hours,
                    notes=notes, created_by=request.user,
                )
            saved += 1

        log_activity(request.user, 'create', 'TemporaryAttendance',
                     description=f'Marked temp attendance for {selected_date} — {saved} recorded, {cleared} cleared',
                     request=request)
        messages.success(request, f'Temporary attendance saved for {selected_date} — {saved} worker(s) recorded.')
        return redirect(f"/temp-attendance/mark/?date={selected_date}")

    workers = TemporaryEmployee.objects.filter(status='active').order_by('name')
    existing_map = {
        a.temp_employee_id: a
        for a in TemporaryAttendance.objects.filter(date=selected_date)
    }
    worker_rows = [{'worker': w, 'att': existing_map.get(w.pk)} for w in workers]

    return render(request, 'attendance_app/temp_attendance_mark.html', {
        'worker_rows': worker_rows,
        'selected_date': selected_date,
        'today': today,
        'prev_date': selected_date - datetime.timedelta(days=1),
        'next_date': selected_date + datetime.timedelta(days=1),
        'marked_count': len(existing_map),
        'total_count': workers.count(),
    })


@login_required
def temp_attendance_list(request):
    if not has_perm(request.user, 'temp_attendance_view'):
        messages.error(request, 'Access denied.')
        return redirect('dashboard')

    records = TemporaryAttendance.objects.select_related('temp_employee').all()
    worker_id = request.GET.get('worker', '')
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')

    if worker_id:
        records = records.filter(temp_employee_id=worker_id)
    if date_from:
        records = records.filter(date__gte=date_from)
    if date_to:
        records = records.filter(date__lte=date_to)

    records = list(records[:500])
    total_hours = sum(float(r.hours_worked) for r in records)
    total_amount = sum(r.amount() or 0 for r in records)

    return render(request, 'attendance_app/temp_attendance_list.html', {
        'records': records,
        'workers': TemporaryEmployee.objects.all().order_by('name'),
        'worker_id': worker_id, 'date_from': date_from, 'date_to': date_to,
        'total_hours': round(total_hours, 2),
        'total_amount': round(total_amount, 2),
        'can_manage': has_perm(request.user, 'temp_attendance_manage'),
    })


@login_required
def temp_attendance_edit(request, pk):
    if not has_perm(request.user, 'temp_attendance_manage'):
        messages.error(request, 'Access denied.')
        return redirect('temp_attendance_list')

    a = get_object_or_404(TemporaryAttendance, pk=pk)
    if request.method == 'POST':
        hours = _f(request, 'hours_worked')
        if not hours or hours <= 0:
            messages.error(request, 'Please enter valid working hours.')
            return render(request, 'attendance_app/temp_attendance_edit.html', {'att': a})

        a.hours_worked = hours
        a.notes = request.POST.get('notes', '').strip()
        a.updated_by = request.user
        a.save()
        log_activity(request.user, 'update', 'TemporaryAttendance', a,
                     f'Updated temp attendance for {a.temp_employee.name} on {a.date}', request)
        messages.success(request, 'Attendance updated.')
        return redirect('temp_attendance_list')

    return render(request, 'attendance_app/temp_attendance_edit.html', {'att': a})


@login_required
def temp_attendance_delete(request, pk):
    if not has_perm(request.user, 'temp_attendance_manage'):
        messages.error(request, 'Access denied.')
        return redirect('temp_attendance_list')

    a = get_object_or_404(TemporaryAttendance, pk=pk)
    if request.method == 'POST':
        desc = f'Deleted temp attendance for {a.temp_employee.name} on {a.date}'
        a.delete()
        log_activity(request.user, 'delete', 'TemporaryAttendance', description=desc, request=request)
        messages.success(request, 'Attendance record deleted.')
    return redirect('temp_attendance_list')


# ─────────────────────────────────────────────────────────────────────────────
# Export — detail rows + auto-summed summary per worker
# ─────────────────────────────────────────────────────────────────────────────

@login_required
def temp_attendance_report(request):
    if not has_perm(request.user, 'temp_attendance_view'):
        messages.error(request, 'Access denied.')
        return redirect('temp_attendance_list')

    today = timezone.localdate()
    return render(request, 'attendance_app/temp_attendance_report.html', {
        'workers': TemporaryEmployee.objects.all().order_by('name'),
        'today': today,
        'month_start': today.replace(day=1),
    })

@login_required
def temp_attendance_export(request):
    if not has_perm(request.user, 'temp_attendance_view'):
        return HttpResponse('Access denied', status=403)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return HttpResponse('openpyxl not installed.', status=500)

    # Accept either multiple ?workers=1&workers=2 (new report page) or a single
    # ?worker=1 (older link) for backward compatibility. Empty = all workers.
    worker_ids = [w for w in request.GET.getlist('workers') if w]
    single_worker = request.GET.get('worker', '')
    if not worker_ids and single_worker:
        worker_ids = [single_worker]

    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')

    # Extra options
    include_notes = request.GET.get('include_notes', '1') == '1'
    include_summary = request.GET.get('include_summary', '1') == '1'
    sort_by = request.GET.get('sort_by', 'worker')  # 'worker' or 'date'

    records = TemporaryAttendance.objects.select_related('temp_employee').all()
    if worker_ids:
        records = records.filter(temp_employee_id__in=worker_ids)
    if date_from:
        records = records.filter(date__gte=date_from)
    if date_to:
        records = records.filter(date__lte=date_to)

    if sort_by == 'date':
        records = list(records.order_by('date', 'temp_employee__name'))
    else:
        records = list(records.order_by('temp_employee__name', 'date'))

    from .models import CompanySettings
    cs = CompanySettings.get_settings()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Temporary Attendance'

    hfill = PatternFill('solid', fgColor='E8650A')
    sub_fill = PatternFill('solid', fgColor='FFF3E8')
    bold_wh = Font(bold=True, color='FFFFFF', size=10)
    bold_or = Font(bold=True, color='E8650A', size=13)
    bold = Font(bold=True, size=10)
    normal = Font(size=10)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(horizontal='left', vertical='center')
    thin = Border(
        left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'),
    )

    period_label = ''
    if date_from and date_to:
        period_label = f' — {date_from} to {date_to}'
    elif date_from:
        period_label = f' — from {date_from}'
    elif date_to:
        period_label = f' — up to {date_to}'

    has_rate = any(r.temp_employee.hourly_rate is not None for r in records)
    headers = ['Date', 'Worker Name', 'Designation', 'Hours Worked']
    if has_rate:
        headers += ['Rate/Hr', 'Amount']
    if include_notes:
        headers.append('Notes')

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws['A1'] = f'{cs.company_name} — Temporary Attendance{period_label}'
    ws['A1'].font = bold_or
    ws['A1'].alignment = center
    ws.row_dimensions[1].height = 22
    ws.append([])

    ws.append(headers)
    hr = ws.max_row
    for ci, h in enumerate(headers, 1):
        c = ws.cell(hr, ci)
        c.value = h; c.font = bold_wh; c.fill = hfill; c.alignment = center; c.border = thin
    ws.row_dimensions[hr].height = 16

    for idx, r in enumerate(records, 1):
        row = [r.date.strftime('%d-%m-%Y'), r.temp_employee.name, r.temp_employee.designation or '', float(r.hours_worked)]
        if has_rate:
            rate = r.temp_employee.hourly_rate
            row += [float(rate) if rate is not None else '', r.amount() if r.amount() is not None else '']
        if include_notes:
            row.append(r.notes or '')
        ws.append(row)
        rr = ws.max_row
        fill = PatternFill('solid', fgColor='FFFFFF') if idx % 2 == 0 else PatternFill('solid', fgColor='FFF8F3')
        for ci in range(1, len(headers) + 1):
            c = ws.cell(rr, ci)
            c.fill = fill; c.border = thin
            c.alignment = left if ci == 2 else center
            c.font = normal

    grand_hours = sum(float(r.hours_worked) for r in records)
    grand_amount = sum(r.amount() or 0 for r in records)

    # ── SUMMARY PER WORKER (auto-totalled) — optional ───────────────────────
    if include_summary:
        ws.append([])
        ws.append([])
        sr = ws.max_row + 2
        ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=len(headers))
        ws.cell(sr, 1, 'SUMMARY (Auto-Calculated)').font = Font(bold=True, size=12, color='E8650A')
        ws.cell(sr, 1).alignment = center
        ws.row_dimensions[sr].height = 18

        sum_headers = ['Worker Name', 'Days Worked', 'Total Hours']
        if has_rate:
            sum_headers += ['Rate/Hr', 'Total Amount']
        ws.append(sum_headers)
        shr = ws.max_row
        for ci, h in enumerate(sum_headers, 1):
            c = ws.cell(shr, ci)
            c.value = h; c.font = bold_wh; c.fill = PatternFill('solid', fgColor='1E293B')
            c.alignment = center; c.border = thin
        ws.row_dimensions[shr].height = 16

        by_worker = defaultdict(list)
        for r in records:
            by_worker[r.temp_employee_id].append(r)

        for temp_id, recs in by_worker.items():
            w = recs[0].temp_employee
            total_h = sum(float(x.hours_worked) for x in recs)
            row = [w.name, len(recs), round(total_h, 2)]
            if has_rate:
                if w.hourly_rate is not None:
                    amt = round(total_h * float(w.hourly_rate), 2)
                    row += [float(w.hourly_rate), amt]
                else:
                    row += ['', '']
            ws.append(row)
            rr = ws.max_row
            for ci in range(1, len(sum_headers) + 1):
                c = ws.cell(rr, ci)
                c.border = thin
                c.alignment = left if ci == 1 else center
                c.font = bold if ci == len(sum_headers) else normal
                c.fill = sub_fill

        ws.append([])
        gtr = ws.max_row + 1
        ws.cell(gtr, 1, 'GRAND TOTAL').font = Font(bold=True, size=11, color='E8650A')
        ws.cell(gtr, 3, round(grand_hours, 2)).font = Font(bold=True, size=11, color='E8650A')
        if has_rate:
            ws.cell(gtr, 5, round(grand_amount, 2)).font = Font(bold=True, size=11, color='E8650A')

    col_widths = [12, 24, 18, 14, 10, 12, 26]
    for ci, w in enumerate(col_widths[:len(headers)], 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    log_activity(request.user, 'export', 'TemporaryAttendance', description='Exported temporary attendance', request=request)
    resp = HttpResponse(buf, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="temporary_attendance_export.xlsx"'
    return resp


@login_required
def temp_employee_detail(request, pk):
    """Dedicated attendance-history view for a single temporary worker."""
    if not has_perm(request.user, 'temp_attendance_view'):
        messages.error(request, 'Access denied.')
        return redirect('temp_employee_list')

    worker = get_object_or_404(TemporaryEmployee, pk=pk)
    records = worker.attendances.all().order_by('-date')

    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    if date_from:
        records = records.filter(date__gte=date_from)
    if date_to:
        records = records.filter(date__lte=date_to)

    records = list(records)
    total_hours = sum(float(r.hours_worked) for r in records)
    total_amount = sum(r.amount() or 0 for r in records)

    return render(request, 'attendance_app/temp_employee_detail.html', {
        'worker': worker,
        'records': records,
        'date_from': date_from, 'date_to': date_to,
        'total_days': len(records),
        'total_hours': round(total_hours, 2),
        'total_amount': round(total_amount, 2),
        'can_manage': has_perm(request.user, 'temp_attendance_manage'),
    })